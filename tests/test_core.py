from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from app import (
    _xlsx_rows_and_fills,
    normalize_matchup,
    parse_date,
    parse_matches,
    parse_start,
    photographer_from_color,
    sheet_export_url,
    unique_destination,
)


HEADERS = [
    "比賽日期", "比賽開始時間", "比賽地點", "年級組別",
    "隊伍名稱/球員姓名/背號", "比賽隊伍名稱", "團體拍攝/指定球員",
]


def test_matchup_normalization():
    cases = {
        "北士足球 vs Ukicker": "北士足球 vs Ukicker",
        "巴格浪U12 VS 明道鋼彈": "巴格浪U12 vs 明道鋼彈",
        "北士足球 vs. Ukicker": "北士足球 vs Ukicker",
        "Faith/永順Fc": "Faith vs 永順Fc",
        "Faith／永順Fc": "Faith vs 永順Fc",
        "巴格浪v新屋": "巴格浪 vs 新屋",
        "猛瑪海王星vsLS紅獅": "猛瑪海王星 vs LS紅獅",
    }
    for source, expected in cases.items():
        assert normalize_matchup(source) == expected


def test_date_and_time_formats():
    assert parse_date("2026年09月12日 禮拜六") == datetime(2026, 9, 12)
    day = datetime(2026, 9, 12)
    assert parse_start(day, "800").strftime("%H:%M") == "08:00"
    assert parse_start(day, "0825").strftime("%H:%M") == "08:25"
    assert parse_start(day, "1330").strftime("%H:%M") == "13:30"
    assert parse_start(day, "920.0").strftime("%H:%M") == "09:20"
    assert parse_start(day, "09:20").strftime("%H:%M") == "09:20"
    assert parse_start(day, 0.5).strftime("%H:%M") == "12:00"


def test_xlsx_color_and_opponent_completion(tmp_path: Path):
    path = tmp_path / "預約表.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(HEADERS)
    sheet.append(["2026年09月12日", "800", "甲", "U10", "大將", "安東國小", "團體"])
    sheet.append(["2026年09月12日", "800", "乙", "U12", "Faith", "Faith/永順Fc", "團體"])
    sheet["D2"].fill = PatternFill("solid", fgColor="FFFF00")
    sheet["D3"].fill = PatternFill("solid", fgColor="00FFFF")
    book.save(path)

    rows, fills = _xlsx_rows_and_fills(path)
    matches = parse_matches(rows, fills)
    got = {(m.photographer, m.matchup) for m in matches}
    assert got == {("植先", "大將 vs 安東國小"), ("植丞", "Faith vs 永順Fc")}


def test_color_detection():
    assert photographer_from_color("FFFF00") == "植先"
    assert photographer_from_color("00FFFF") == "植丞"
    assert photographer_from_color("") == "待確認"


def test_google_export_url_keeps_xlsx():
    url = "https://docs.google.com/spreadsheets/d/abc123/edit#gid=99"
    assert sheet_export_url(url, "xlsx") == "https://docs.google.com/spreadsheets/d/abc123/export?format=xlsx"


def test_duplicate_filename_is_not_overwritten(tmp_path: Path):
    original = tmp_path / "DSC0001.JPG"
    original.write_bytes(b"first")
    assert unique_destination(original).name == "DSC0001_2.JPG"
